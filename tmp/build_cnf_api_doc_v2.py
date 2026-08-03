from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(r"C:\Users\sidmy\Downloads\shree-cement-cnf-agent-cowork-20260729T062834Z-1-001\shree-cement-cnf-agent-cowork")
SOURCE = WORKSPACE / "deliverables" / "SHREE_CNF_SAP_API_DOCUMENTATION_PRE_DEV_V2.md"
OUTPUT = WORKSPACE / "deliverables" / "SHREE_CNF_SAP_API_DOCUMENTATION_PRE_DEV_V2.docx"

INK = "1F2937"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
GRID = "C9D2DC"
CAUTION_FILL = "FFF4D6"
CAUTION_TEXT = "7A5A00"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_borders(table, color=GRID, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def paragraph_border_bottom(paragraph, color=BLUE, size="12", space="6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_field(paragraph, instruction, fallback):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, size=9, color=MID_GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(INK)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.paragraph_format.space_after = Pt(16)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def add_inline_markdown(paragraph, text, default_size=11, default_color=INK):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=default_size, color=default_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=default_color, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=max(8.5, default_size - 1), color=DARK_BLUE)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=default_size, color=default_color)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def widths_for(columns):
    choices = {
        2: [2400, 6960],
        3: [1700, 3500, 4160],
        4: [1200, 3100, 1800, 3260],
        5: [900, 2200, 1100, 1600, 3560],
        6: [750, 1500, 950, 1300, 2300, 2560],
    }
    return choices.get(columns, [9360 // columns] * columns)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    set_table_geometry(table, widths_for(cols))
    set_borders(table)
    repeat_header(table.rows[0])
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx in range(cols):
            value = row_data[c_idx] if c_idx < len(row_data) else ""
            cell = row.cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if c_idx == 0 or value.replace("+", "").replace("-", "").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_markdown(p, value, default_size=8.5 if cols >= 5 else 9)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_borders(table, color="E3C56E", size="6")
    repeat_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CAUTION_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    add_inline_markdown(p, text, default_size=10.5, default_color=CAUTION_TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run("Shree Cement C&F Agent Interface | SAP API Pre-DEV")
    set_run_font(run, size=8.5, color=MID_GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("CNF-SAP-API-DOC-002 v2.0  |  ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    add_field(fp, "PAGE", "1")

    # Cover / memo masthead.
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Shree Cement C&F Agent Interface")
    p2 = doc.add_paragraph(style="Subtitle")
    p2.add_run("SAP API Documentation and Pre-Development Proposal")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(16)
    paragraph_border_bottom(rule, color=BLUE, size="16", space="2")

    metadata = [
        ("Document ID", "CNF-SAP-API-DOC-002"),
        ("Version", "2.0"),
        ("Date", "30 July 2026"),
        ("Audience", "SAP, Integration, Commerce, Datasphere, Basis, Security, Tax and functional teams"),
        ("Owner", "SAP backend integration workstream"),
        ("Status", "PRE-DEV WORKING DOCUMENT - NOT AN APPROVED BUILD SPECIFICATION"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label + ": ")
        set_run_font(r, size=10.5, color=INK, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    callout_line = next(line[2:].strip() for line in lines if line.startswith("> "))
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_callout(doc, callout_line)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Prepared from the latest five delivery artifacts and reconciled project context")
    set_run_font(r, size=10, color=MID_GRAY, italic=True)
    doc.add_page_break()

    # Skip opening title, metadata and cover callout from Markdown.
    start = 0
    while start < len(lines) and not lines[start].startswith("# 1. "):
        start += 1

    i = start
    first_h1 = True
    pending_para = []

    def flush_para():
        nonlocal pending_para
        if pending_para:
            text = " ".join(part.strip() for part in pending_para).strip()
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.widow_control = True
                add_inline_markdown(p, text)
            pending_para = []

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            flush_para()
            i += 1
            continue
        if line.startswith("|"):
            flush_para()
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            flush_para()
            if not first_h1:
                doc.add_page_break()
            first_h1 = False
            p = doc.add_paragraph(style="Heading 1")
            add_inline_markdown(p, line[2:], default_size=16, default_color=BLUE)
            i += 1
            continue
        if line.startswith("## "):
            flush_para()
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, line[3:], default_size=13, default_color=BLUE)
            i += 1
            continue
        if line.startswith("### "):
            flush_para()
            p = doc.add_paragraph(style="Heading 3")
            add_inline_markdown(p, line[4:], default_size=12, default_color=DARK_BLUE)
            i += 1
            continue
        if line.startswith("> "):
            flush_para()
            add_callout(doc, line[2:])
            i += 1
            continue
        if re.match(r"^- ", line):
            flush_para()
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, line[2:])
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            flush_para()
            p = doc.add_paragraph(style="List Number")
            add_inline_markdown(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        pending_para.append(line)
        i += 1
    flush_para()

    # Document properties.
    props = doc.core_properties
    props.title = "Shree Cement C&F Agent Interface - SAP API Documentation and Pre-Development Proposal"
    props.subject = "Consolidated pre-development SAP API documentation"
    props.author = "SAP backend integration workstream"
    props.keywords = "Shree Cement, C&F, SAP, API, S/4HANA, Integration Suite, pre-development"
    props.comments = "Working proposal. SAP development access pending."
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
